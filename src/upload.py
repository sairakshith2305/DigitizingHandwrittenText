import os
from datetime import datetime
from flask import Flask, request, render_template, send_from_directory, send_file, url_for, redirect
from docx import Document
from fpdf import FPDF
from main import infer_by_web

__author__ = 'mini'

app = Flask(__name__)

APP_ROOT = os.path.dirname(os.path.abspath(__file__)) # project abs path

class setandgetvalue:
	def __init__(self,value):
		self.value = value
	def setvalue(self,value):
		self.value = value
	def getvalue(self):
		return self.value

a = setandgetvalue("") 
	

@app.route("/")
def index():
    return render_template("index.html")


@app.route("/upload_page", methods=["GET"])
def upload_page():
    return render_template("upload.html")


@app.route("/upload", methods=["POST"])
def upload():
    # folder_name = request.form['uploads']
    target = os.path.join(APP_ROOT, 'static/')
    print(target)
    if not os.path.isdir(target):
        os.mkdir(target)
    print(request.files.getlist("file"))
    option = request.form.get('optionsPrediction')
    print("Selected Option:: {}".format(option))
    for upload in request.files.getlist("file"):
        print(upload)
        print("{} is the file name".format(upload.filename))
        filename = upload.filename
        # This is to verify files are supported
        ext = os.path.splitext(filename)[1]
        if (ext == ".jpg") or (ext == ".png"):
            print("File supported moving on...")
        else:
            render_template("Error.html", message="Files uploaded are not supported...")
        savefname = datetime.now().strftime('%Y-%m-%d_%H_%M_%S') + "."+ext
        destination = "/".join([target, savefname])
        print("Accept incoming file:", filename)
        print("Save it to:", destination)
        upload.save(destination)
        x = predict_image(destination, option)
	#a.setvalue(x)
        print("Prediction: ", x)
    # return send_from_directory("images", filename, as_attachment=True)
    a.setvalue(x)
    return render_template("complete.html", image_name=savefname, result=x)

@app.route('/download', methods=['POST'])
def save_file():
    option = request.form['action']
    #texts = request.form.getlist("message[]")
    texts = a.getvalue()
    print(texts)
    if option == "word":
      document = Document()
      #texts = texts.split("\n")
      #r_text = []
      #for line in texts:
      #  if not line.isspace():
      #  if line != "\n":
      #          r_text.append(line)
      #texts = "".join(r_text[1:])
      for text in texts:
        document.add_paragraph(text.replace("\n", ""))
      document.save('extracted_text.docx')
      return send_file("extracted_text.docx",as_attachment=True, cache_timeout=0)

    elif option == "pdf":
      texts = "".join(texts)
      pdf = FPDF()
      pdf.alias_nb_pages()
      pdf.add_page()
      pdf.set_font('Times', '', 10)
      for text in texts.split("\n"):
        pdf.cell(0, 8, text.encode('utf-8').decode('latin-1'), 0, 1)
      pdf.output('extracted_text.pdf', 'F')
      return send_file("extracted_text.pdf",as_attachment=True, cache_timeout=0)

    elif option == "txt":
      with open("extracted_text.txt", "w") as f:
        for text in texts:
          f.write(text.replace("\n", ""))
      return send_file("extracted_text.txt",as_attachment=True, cache_timeout=0)

    return ""



def predict_image(path, type):
    print(path)
    return infer_by_web(path, type)


if __name__ == "__main__":
    app.run(port=4555, debug=True)
